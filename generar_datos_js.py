# -*- coding: utf-8 -*-
"""Genera `datos.js`: copia de los 4 Excels como JSON embebido en la pagina.

POR QUE: abriendo index.html con doble clic (file://) el navegador BLOQUEA la
lectura de los .xlsx con fetch() (no se puede leer un archivo local por
AJAX). Con este respaldo, la pagina tambien funciona en file:// como PRUEBA:
script.js intenta leer el Excel y, si falla, usa window.LYGANDO_DATOS.

COMO:  python3 generar_datos_js.py   (volver a correrlo tras registrar ofertas
nuevas, cambiar un Excel o re-scrapear).
"""
import json
import os
import datetime
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
ARCHIVOS = [
    "ofertas.xlsx",
    "convocatorias.xlsx",
    "convocatoriasdetrabajo.xlsx",
    "cursos.xlsx",
]
SALIDA = os.path.join(BASE, "datos.js")


def limpio(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, bool):
        return "si" if v else "no"
    return v


def filas_de(archivo):
    wb = openpyxl.load_workbook(os.path.join(BASE, archivo), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = list(ws.iter_rows(values_only=True))
    wb.close()
    if not filas:
        return []
    header = [str(h).strip() if h is not None else "" for h in filas[0]]
    salida = []
    for f in filas[1:]:
        if f is None or f[0] is None or str(f[0]).strip() == "":
            continue
        d = {}
        for i, col in enumerate(header):
            if not col:
                continue
            d[col] = limpio(f[i] if i < len(f) else "")
        salida.append(d)
    return salida


def main():
    datos = {}
    for a in ARCHIVOS:
        if not os.path.exists(os.path.join(BASE, a)):
            print("FALTA", a, "-> se omite")
            continue
        filas = filas_de(a)
        datos[a] = filas
        print(f"{a}: {len(filas)} filas")
    with open(SALIDA, "w", encoding="utf-8") as fh:
        fh.write("// Generado por generar_datos_js.py - NO EDITAR A MANO.\n")
        fh.write("// Respaldo de los Excels para que index.html funcione tambien en file://.\n")
        fh.write("window.LYGANDO_DATOS = ")
        json.dump(datos, fh, ensure_ascii=False, separators=(",", ":"))
        fh.write(";\n")
    print("Escrito", SALIDA, f"({os.path.getsize(SALIDA)//1024} KB)")


if __name__ == "__main__":
    main()
